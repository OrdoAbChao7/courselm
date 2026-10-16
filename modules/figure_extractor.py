"""High-fidelity extraction and cropping engine for CourseLM Source Figures.

Adheres strictly to the provenance principles:
- Only real images extracted or cropped from source documents.
- No AI redrawing, no generative inpainting, no hallucinatory rendering.
- Full provenance metadata: source file, SHA-256, page/slide number, coordinates.
- High-resolution rasterization (300 DPI) for vector circuits, mechanics diagrams,
  coordinate plots, and scanned pages.
- Deduplication by image SHA-256 with multiple occurrence records preserved.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable, List, Optional, Tuple

import numpy as np
import pypdfium2 as pdfium
from loguru import logger
from PIL import Image

from modules.figure_models import (
    BoundingBox,
    CoordinateSystem,
    ExtractionType,
    FigureAsset,
    FigureManifest,
    ReviewStatus,
    SourceOccurrence,
)


def compute_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def compute_file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def is_blank_image(image: Image.Image, threshold_std: float = 2.0) -> bool:
    """Check if an image is blank (pure white, pure black, or nearly uniform).

    Inspired by marker's blank check, optimized with numpy std deviation.
    """
    if image.width < 8 or image.height < 8:
        return True
    # Analyze visible pixels, without changing the stored original image.
    rgba = image.convert('RGBA')
    visible = Image.new('RGBA', rgba.size, 'white')
    visible.alpha_composite(rgba)
    visible = visible.convert('RGB')
    visible.thumbnail((96, 96))
    arr = np.asarray(visible, dtype=float)
    if arr.size == 0:
        return True
    std = float(np.std(arr))
    if float(np.max(np.std(arr, axis=(0, 1)))) < threshold_std:
        return True
    # A smooth background gradient has substantial global variance but almost no
    # residual after a low-order surface fit. It is not evidence of a diagram.
    height, width = arr.shape[:2]
    y, x = np.mgrid[-1:1:complex(height), -1:1:complex(width)]
    design = np.column_stack([np.ones(x.size), x.ravel(), y.ravel(),
                              (x*y).ravel(), (x*x).ravel(), (y*y).ravel()])
    colors = arr.reshape(-1, 3)
    coefficients, *_ = np.linalg.lstsq(design, colors, rcond=None)
    residual = colors - design @ coefficients
    return float(np.sqrt(np.mean(residual ** 2))) < 1.5


def _extract_text_in_rect(
    textpage: pdfium.PdfTextPage,
    rect: Tuple[float, float, float, float], # (left, bottom, right, top) in points
    padding: float = 30.0,
) -> str:
    """Extract text lines within or immediately adjacent to a rectangular region in points."""
    try:
        left, bottom, right, top = rect
        expanded_rect = (
            max(0.0, left - padding),
            max(0.0, bottom - padding),
            right + padding,
            top + padding,
        )
        # pypdfium2 get_text_bounded takes (left, top, right, bottom) where top > bottom
        # or (left, bottom, right, top). Let's use get_text_bounded:
        return textpage.get_text_bounded(
            expanded_rect[0], expanded_rect[1], expanded_rect[2], expanded_rect[3]
        ).strip()
    except Exception:
        return ""


class FigureExtractor:
    """Extracts and manages source figures from PDF, PPTX, DOCX, and images."""

    def __init__(self, target_dpi: int = 300):
        self.target_dpi = target_dpi
        self.render_scale = target_dpi / 72.0

    def extract_from_pdf(
        self,
        pdf_path: Path,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        manifest: FigureManifest,
        max_pages: Optional[int] = None,
        extract_vectors: bool = True,
    ) -> List[FigureAsset]:
        """Extract embedded figures and vector graphics clusters from a PDF document."""
        task_figures_dir.mkdir(parents=True, exist_ok=True)
        source_sha = compute_file_sha256(pdf_path)
        doc = pdfium.PdfDocument(pdf_path)
        total_pages = len(doc)
        pages_to_process = min(total_pages, max_pages) if max_pages else total_pages
        extracted: List[FigureAsset] = []

        # Track existing hashes in this task for deduplication
        hash_to_asset: dict[str, FigureAsset] = {
            a.image_sha256: a for a in manifest.figures.values()
        }

        for page_idx in range(pages_to_process):
            page_num = page_idx + 1
            page = doc[page_idx]
            page_w, page_h = page.get_size()
            rotation = page.get_rotation()
            textpage = page.get_textpage()

            # 1. Enumerate page objects to find embedded images
            try:
                objs = list(page.get_objects())
            except Exception as e:
                logger.warning(f"Failed to get objects on {pdf_path.name} p.{page_num}: {e}")
                objs = []

            image_objs = [o for o in objs if o.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE]
            path_objs = [o for o in objs if o.type == pdfium.raw.FPDF_PAGEOBJ_PATH]

            # A. Process embedded raster images
            for img_idx, img_obj in enumerate(image_objs):
                try:
                    bounds = img_obj.get_bounds() # (left, bottom, right, top)
                    left, bottom, right, top = bounds
                    box_w = right - left
                    box_h = top - bottom

                    # Filter out tiny icon/spacer images (< 20 pt in either dimension)
                    if box_w < 20 or box_h < 20:
                        continue

                    # Try to extract bitmap
                    bitmap = img_obj.get_bitmap()
                    pil_img = bitmap.to_pil()
                    if is_blank_image(pil_img):
                        continue

                    # Encode to PNG bytes
                    buf = io.BytesIO()
                    pil_img.save(buf, format="PNG", optimize=True)
                    img_bytes = buf.getvalue()
                    img_sha = compute_sha256(img_bytes)

                    # Normalized bbox (top-left origin 0..1 for standard UI coordinates)
                    norm_bbox = BoundingBox(
                        x0=max(0.0, min(1.0, left / page_w)),
                        y0=max(0.0, min(1.0, (page_h - top) / page_h)),
                        x1=max(0.0, min(1.0, right / page_w)),
                        y1=max(0.0, min(1.0, (page_h - bottom) / page_h)),
                        coord_system=CoordinateSystem.NORMALIZED,
                    )

                    occ = SourceOccurrence(
                        source_file=rel_source_path,
                        source_sha256=source_sha,
                        page_index=page_num,
                        bbox=norm_bbox,
                    )

                    # Check deduplication
                    if img_sha in hash_to_asset:
                        existing = hash_to_asset[img_sha]
                        existing.add_occurrence(occ)
                        continue

                    adjacent = _extract_text_in_rect(textpage, bounds, padding=40.0)
                    asset_id = f"fig_{img_sha[:12]}"
                    image_filename = f"{asset_id}.png"
                    image_disk_path = task_figures_dir / image_filename
                    image_disk_path.write_bytes(img_bytes)

                    caption = self._infer_caption(adjacent, page_num)

                    asset = FigureAsset(
                        asset_id=asset_id,
                        task_id=task_id,
                        source_file=rel_source_path,
                        source_sha256=source_sha,
                        page_index=page_num,
                        extraction_type=ExtractionType.EMBEDDED,
                        bbox=norm_bbox,
                        page_rotation=rotation,
                        image_rel_path=f"figures/{image_filename}",
                        image_sha256=img_sha,
                        width=pil_img.width,
                        height=pil_img.height,
                        dpi=self.target_dpi,
                        adjacent_text=adjacent[:500],
                        caption=caption,
                        status=ReviewStatus.PENDING,
                        occurrences=[occ],
                        metadata={"format": "png", "page_size": [page_w, page_h]},
                    )
                    manifest.add(asset)
                    hash_to_asset[img_sha] = asset
                    extracted.append(asset)
                except Exception as e:
                    logger.debug(f"Failed extracting image object {img_idx} on p.{page_num}: {e}")

            # B. Vector / Graphical Cluster Detection (Circuits, Force Diagrams, Plots)
            if extract_vectors and path_objs:
                vector_clusters = self._cluster_vector_paths(path_objs, page_w, page_h)
                for c_idx, cluster_box in enumerate(vector_clusters):
                    try:
                        # cluster_box is (left, bottom, right, top) in points
                        c_left, c_bottom, c_right, c_top = cluster_box
                        c_w = c_right - c_left
                        c_h = c_top - c_bottom
                        # Must be meaningful figure size
                        if c_w < 40 or c_h < 30:
                            continue

                        # Check if this cluster overlaps significantly with an already extracted embedded image
                        overlap = False
                        for img_obj in image_objs:
                            ib = img_obj.get_bounds()
                            if self._boxes_intersect(cluster_box, ib, threshold=0.7):
                                overlap = True
                                break
                        if overlap:
                            continue

                        # Render and crop page region
                        norm_bbox = BoundingBox(
                            x0=max(0.0, min(1.0, c_left / page_w)),
                            y0=max(0.0, min(1.0, (page_h - c_top) / page_h)),
                            x1=max(0.0, min(1.0, c_right / page_w)),
                            y1=max(0.0, min(1.0, (page_h - c_bottom) / page_h)),
                            coord_system=CoordinateSystem.NORMALIZED,
                        )

                        crop_asset = self.crop_pdf_region(
                            pdf_path=pdf_path,
                            page_index=page_num,
                            bbox=norm_bbox,
                            task_figures_dir=task_figures_dir,
                            task_id=task_id,
                            rel_source_path=rel_source_path,
                            manifest=manifest,
                            dpi=self.target_dpi,
                            extraction_type=ExtractionType.CROPPED,
                            status=ReviewStatus.PENDING,
                        )
                        if crop_asset:
                            extracted.append(crop_asset)
                            hash_to_asset[crop_asset.image_sha256] = crop_asset
                    except Exception as e:
                        logger.debug(f"Vector cluster extraction failed on p.{page_num}: {e}")

        return extracted

    def crop_pdf_region(
        self,
        pdf_path: Path,
        page_index: int,
        bbox: BoundingBox,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        manifest: FigureManifest,
        dpi: int = 300,
        caption: str = "",
        status: ReviewStatus = ReviewStatus.CONFIRMED,
        extraction_type: ExtractionType = ExtractionType.CROPPED,
    ) -> Optional[FigureAsset]:
        """Crop a specified bounding box from a PDF page at high DPI."""
        task_figures_dir.mkdir(parents=True, exist_ok=True)
        source_sha = compute_file_sha256(pdf_path)
        doc = pdfium.PdfDocument(pdf_path)
        if page_index <= 0:
            page_index = 1
        if page_index > len(doc):
            raise ValueError(f"Page index {page_index} out of range (1..{len(doc)})")

        page = doc[page_index - 1]
        page_w, page_h = page.get_size()
        rotation = page.get_rotation()
        if bbox.x1 > 1.0 or bbox.y1 > 1.0 or bbox.x0 > 1.0 or bbox.y0 > 1.0:
            norm_bbox = BoundingBox(
                x0=max(0.0, min(1.0, bbox.x0 / page_w)),
                y0=max(0.0, min(1.0, bbox.y0 / page_h)),
                x1=max(0.0, min(1.0, bbox.x1 / page_w)),
                y1=max(0.0, min(1.0, bbox.y1 / page_h)),
                coord_system=CoordinateSystem.NORMALIZED,
            )
        else:
            norm_bbox = bbox.to_normalized(page_w, page_h)

        scale = dpi / 72.0
        rendered_page = page.render(scale=scale).to_pil()
        rw, rh = rendered_page.size

        # Pixel bounds in rendered page (top-left origin)
        px0 = max(0, min(rw, int(norm_bbox.x0 * rw)))
        py0 = max(0, min(rh, int(norm_bbox.y0 * rh)))
        px1 = max(0, min(rw, int(norm_bbox.x1 * rw)))
        py1 = max(0, min(rh, int(norm_bbox.y1 * rh)))

        # Ensure valid area
        if px1 - px0 < 4 or py1 - py0 < 4:
            logger.warning("Cropping region too small")
            return None

        cropped = rendered_page.crop((px0, py0, px1, py1))
        if is_blank_image(cropped):
            logger.debug("Cropped image is blank, discarding")
            return None

        buf = io.BytesIO()
        cropped.save(buf, format="PNG", optimize=True)
        img_bytes = buf.getvalue()
        img_sha = compute_sha256(img_bytes)

        occ = SourceOccurrence(
            source_file=rel_source_path,
            source_sha256=source_sha,
            page_index=page_index,
            bbox=norm_bbox,
        )

        # Check deduplication
        existing = manifest.get(f"fig_{img_sha[:12]}")
        if existing:
            existing.add_occurrence(occ)
            if caption:
                existing.caption = caption
            # Rescanning cannot undo an explicit rejection or downgrade approval.
            if existing.status == ReviewStatus.PENDING or status == ReviewStatus.REJECTED:
                existing.status = status
            return existing

        textpage = page.get_textpage()
        # Convert norm_bbox back to PDF points (left, bottom, right, top)
        pt_left = norm_bbox.x0 * page_w
        pt_right = norm_bbox.x1 * page_w
        pt_top = page_h - (norm_bbox.y0 * page_h)
        pt_bottom = page_h - (norm_bbox.y1 * page_h)
        adjacent = _extract_text_in_rect(textpage, (pt_left, pt_bottom, pt_right, pt_top), padding=50.0)

        if not caption:
            caption = self._infer_caption(adjacent, page_index)

        asset_id = f"fig_{img_sha[:12]}"
        image_filename = f"{asset_id}.png"
        image_disk_path = task_figures_dir / image_filename
        image_disk_path.write_bytes(img_bytes)

        asset = FigureAsset(
            asset_id=asset_id,
            task_id=task_id,
            source_file=rel_source_path,
            source_sha256=source_sha,
            page_index=page_index,
            extraction_type=extraction_type,
            bbox=norm_bbox,
            page_rotation=rotation,
            image_rel_path=f"figures/{image_filename}",
            image_sha256=img_sha,
            width=cropped.width,
            height=cropped.height,
            dpi=dpi,
            adjacent_text=adjacent[:500],
            caption=caption,
            status=status,
            occurrences=[occ],
            metadata={"format": "png", "crop_px": [px0, py0, px1, py1]},
        )
        manifest.add(asset)
        return asset

    def extract_from_pptx(
        self,
        pptx_path: Path,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        manifest: FigureManifest,
    ) -> List[FigureAsset]:
        """Extract pictures and shapes from a PowerPoint presentation."""
        try:
            import pptx
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError:
            logger.warning("python-pptx not available for PPTX extraction")
            return []

        task_figures_dir.mkdir(parents=True, exist_ok=True)
        source_sha = compute_file_sha256(pptx_path)
        prs = pptx.Presentation(pptx_path)
        extracted: List[FigureAsset] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            # Extract slide text context
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text:
                    slide_texts.append(shape.text_frame.text.strip())
            context_text = " \n".join(slide_texts)[:500]

            slide_pic_count = 0
            for shape_idx, shape in enumerate(slide.shapes):
                # Direct picture
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                    try:
                        img_bytes = shape.image.blob
                        pil_img = Image.open(io.BytesIO(img_bytes))
                        if is_blank_image(pil_img):
                            continue

                        buf = io.BytesIO()
                        pil_img.save(buf, format="PNG")
                        clean_bytes = buf.getvalue()
                        img_sha = compute_sha256(clean_bytes)

                        occ = SourceOccurrence(
                            source_file=rel_source_path,
                            source_sha256=source_sha,
                            page_index=slide_num,
                            bbox=None,
                        )

                        existing = manifest.get(f"fig_{img_sha[:12]}")
                        if existing:
                            existing.add_occurrence(occ)
                            slide_pic_count += 1
                            continue

                        asset_id = f"fig_{img_sha[:12]}"
                        image_filename = f"{asset_id}.png"
                        (task_figures_dir / image_filename).write_bytes(clean_bytes)

                        caption = self._infer_caption(context_text, slide_num) or f"幻灯片 {slide_num} 插图"

                        asset = FigureAsset(
                            asset_id=asset_id,
                            task_id=task_id,
                            source_file=rel_source_path,
                            source_sha256=source_sha,
                            page_index=slide_num,
                            extraction_type=ExtractionType.EMBEDDED,
                            bbox=None,
                            page_rotation=0,
                            image_rel_path=f"figures/{image_filename}",
                            image_sha256=img_sha,
                            width=pil_img.width,
                            height=pil_img.height,
                            dpi=self.target_dpi,
                            adjacent_text=context_text,
                            caption=caption,
                            status=ReviewStatus.PENDING,
                            occurrences=[occ],
                            metadata={"shape_name": shape.name},
                        )
                        manifest.add(asset)
                        extracted.append(asset)
                        slide_pic_count += 1
                    except Exception as e:
                        logger.debug(f"Failed extracting PPTX picture on slide {slide_num}: {e}")

            # If no pictures were extracted but the slide contains graphic shapes or diagrams:
            has_graphic_shapes = any(
                getattr(s, "shape_type", None) in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM)
                for s in slide.shapes
            )
            if slide_pic_count == 0 and has_graphic_shapes:
                slide_asset = self._export_pptx_slide_com(
                    pptx_path=pptx_path,
                    slide_num=slide_num,
                    task_figures_dir=task_figures_dir,
                    task_id=task_id,
                    rel_source_path=rel_source_path,
                    source_sha=source_sha,
                    context_text=context_text,
                    manifest=manifest,
                )
                if slide_asset:
                    extracted.append(slide_asset)

        return extracted

    def _export_pptx_slide_com(
        self,
        pptx_path: Path,
        slide_num: int,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        source_sha: str,
        context_text: str,
        manifest: FigureManifest,
    ) -> Optional[FigureAsset]:
        """Render a vector/diagram slide via PowerPoint COM automation in an isolated process."""
        import subprocess
        import sys

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp_png = Path(tmp.name)

        script = (
            "import win32com.client, sys\n"
            "app = win32com.client.Dispatch('PowerPoint.Application')\n"
            "pres = None\n"
            "try:\n"
            f"    pres = app.Presentations.Open(sys.argv[1], ReadOnly=True, Untitled=False, WithWindow=False)\n"
            f"    pres.Slides({slide_num}).Export(sys.argv[2], 'PNG', 1920, 1080)\n"
            "finally:\n"
            "    if pres:\n"
            "        try:\n"
            "            pres.Close()\n"
            "        except Exception:\n"
            "            pass\n"
            "    try:\n"
            "        app.Quit()\n"
            "    except Exception:\n"
            "        pass\n"
        )
        try:
            res = subprocess.run(
                [sys.executable, "-c", script, str(pptx_path.resolve()), str(tmp_png)],
                capture_output=True,
                timeout=20,
            )
            if res.returncode == 0 and tmp_png.is_file() and tmp_png.stat().st_size > 0:
                img_bytes = tmp_png.read_bytes()
                pil_img = Image.open(io.BytesIO(img_bytes))
                if not is_blank_image(pil_img):
                    img_sha = compute_sha256(img_bytes)
                    asset_id = f"fig_{img_sha[:12]}"
                    image_filename = f"{asset_id}.png"
                    (task_figures_dir / image_filename).write_bytes(img_bytes)

                    occ = SourceOccurrence(
                        source_file=rel_source_path,
                        source_sha256=source_sha,
                        page_index=slide_num,
                        bbox=None,
                    )
                    caption = self._infer_caption(context_text, slide_num) or f"幻灯片 {slide_num} 结构示意图"
                    asset = FigureAsset(
                        asset_id=asset_id,
                        task_id=task_id,
                        source_file=rel_source_path,
                        source_sha256=source_sha,
                        page_index=slide_num,
                        extraction_type=ExtractionType.CROPPED,
                        bbox=None,
                        page_rotation=0,
                        image_rel_path=f"figures/{image_filename}",
                        image_sha256=img_sha,
                        width=pil_img.width,
                        height=pil_img.height,
                        dpi=self.target_dpi,
                        adjacent_text=context_text,
                        caption=caption,
                        status=ReviewStatus.PENDING,
                        occurrences=[occ],
                        metadata={"rendered_via": "PowerPoint.Application"},
                    )
                    manifest.add(asset)
                    return asset
        except Exception as e:
            logger.debug(f"PowerPoint COM export skipped: {e}")
        finally:
            tmp_png.unlink(missing_ok=True)
        return None

    def extract_from_docx(
        self,
        docx_path: Path,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        manifest: FigureManifest,
    ) -> List[FigureAsset]:
        """Extract embedded images from Word DOCX document."""
        task_figures_dir.mkdir(parents=True, exist_ok=True)
        source_sha = compute_file_sha256(docx_path)
        extracted: List[FigureAsset] = []

        try:
            with zipfile.ZipFile(docx_path, "r") as zf:
                media_files = [n for n in zf.namelist() if n.startswith("word/media/")]
                for idx, media_name in enumerate(media_files):
                    img_data = zf.read(media_name)
                    try:
                        pil_img = Image.open(io.BytesIO(img_data))
                        if is_blank_image(pil_img):
                            continue

                        buf = io.BytesIO()
                        pil_img.save(buf, format="PNG")
                        clean_bytes = buf.getvalue()
                        img_sha = compute_sha256(clean_bytes)

                        occ = SourceOccurrence(
                            source_file=rel_source_path,
                            source_sha256=source_sha,
                            page_index=idx + 1,
                            bbox=None,
                        )

                        existing = manifest.get(f"fig_{img_sha[:12]}")
                        if existing:
                            existing.add_occurrence(occ)
                            continue

                        asset_id = f"fig_{img_sha[:12]}"
                        image_filename = f"{asset_id}.png"
                        (task_figures_dir / image_filename).write_bytes(clean_bytes)

                        asset = FigureAsset(
                            asset_id=asset_id,
                            task_id=task_id,
                            source_file=rel_source_path,
                            source_sha256=source_sha,
                            page_index=idx + 1,
                            extraction_type=ExtractionType.EMBEDDED,
                            bbox=None,
                            page_rotation=0,
                            image_rel_path=f"figures/{image_filename}",
                            image_sha256=img_sha,
                            width=pil_img.width,
                            height=pil_img.height,
                            dpi=self.target_dpi,
                            adjacent_text="",
                            caption=f"文档插图 {idx + 1}",
                            status=ReviewStatus.PENDING,
                            occurrences=[occ],
                            metadata={"media_part": media_name},
                        )
                        manifest.add(asset)
                        extracted.append(asset)
                    except Exception as e:
                        logger.debug(f"Failed decoding docx media {media_name}: {e}")
        except Exception as e:
            logger.warning(f"Failed reading docx {docx_path.name}: {e}")

        return extracted

    def extract_from_image(
        self,
        image_path: Path,
        task_figures_dir: Path,
        task_id: str,
        rel_source_path: str,
        manifest: FigureManifest,
    ) -> Optional[FigureAsset]:
        """Ingest a standalone image file (PNG, JPG, etc.)."""
        task_figures_dir.mkdir(parents=True, exist_ok=True)
        source_sha = compute_file_sha256(image_path)
        try:
            with Image.open(image_path) as pil_img:
                if is_blank_image(pil_img):
                    return None
                buf = io.BytesIO()
                pil_img.save(buf, format="PNG")
                clean_bytes = buf.getvalue()
                img_sha = compute_sha256(clean_bytes)

                occ = SourceOccurrence(
                    source_file=rel_source_path,
                    source_sha256=source_sha,
                    page_index=1,
                    bbox=None,
                )

                existing = manifest.get(f"fig_{img_sha[:12]}")
                if existing:
                    existing.add_occurrence(occ)
                    return existing

                asset_id = f"fig_{img_sha[:12]}"
                image_filename = f"{asset_id}.png"
                (task_figures_dir / image_filename).write_bytes(clean_bytes)

                caption = image_path.stem

                asset = FigureAsset(
                    asset_id=asset_id,
                    task_id=task_id,
                    source_file=rel_source_path,
                    source_sha256=source_sha,
                    page_index=1,
                    extraction_type=ExtractionType.STANDALONE,
                    bbox=None,
                    page_rotation=0,
                    image_rel_path=f"figures/{image_filename}",
                    image_sha256=img_sha,
                    width=pil_img.width,
                    height=pil_img.height,
                    dpi=self.target_dpi,
                    adjacent_text="",
                    caption=caption,
                    status=ReviewStatus.PENDING,
                    occurrences=[occ],
                    metadata={"original_name": image_path.name},
                )
                manifest.add(asset)
                return asset
        except Exception as e:
            logger.warning(f"Failed ingesting standalone image {image_path.name}: {e}")
            return None

    def scan_and_extract_directory(
        self,
        source_dir: Path,
        task_output_dir: Path,
        task_id: str,
        manifest: FigureManifest,
        progress_callback: Optional[Callable[[str, int, int], None]] = None,
    ) -> FigureManifest:
        """Scan a source materials directory and automatically build the figure asset library."""
        task_figures_dir = task_output_dir / "figures"
        task_figures_dir.mkdir(parents=True, exist_ok=True)

        supported_exts = {".pdf", ".pptx", ".ppt", ".docx", ".doc", ".png", ".jpg", ".jpeg", ".webp", ".bmp"}
        files: List[Path] = []
        for p in source_dir.rglob("*"):
            if p.is_file() and p.suffix.lower() in supported_exts:
                files.append(p)

        total_files = len(files)
        logger.info(f"Scanning source directory {source_dir}: found {total_files} files with figures")

        for idx, file_path in enumerate(files):
            rel_path = file_path.relative_to(source_dir).as_posix()
            ext = file_path.suffix.lower()
            if progress_callback:
                progress_callback(rel_path, idx + 1, total_files)

            try:
                if ext == ".pdf":
                    self.extract_from_pdf(
                        pdf_path=file_path,
                        task_figures_dir=task_figures_dir,
                        task_id=task_id,
                        rel_source_path=rel_path,
                        manifest=manifest,
                    )
                elif ext in (".pptx", ".ppt"):
                    self.extract_from_pptx(
                        pptx_path=file_path,
                        task_figures_dir=task_figures_dir,
                        task_id=task_id,
                        rel_source_path=rel_path,
                        manifest=manifest,
                    )
                elif ext in (".docx", ".doc"):
                    self.extract_from_docx(
                        docx_path=file_path,
                        task_figures_dir=task_figures_dir,
                        task_id=task_id,
                        rel_source_path=rel_path,
                        manifest=manifest,
                    )
                elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
                    self.extract_from_image(
                        image_path=file_path,
                        task_figures_dir=task_figures_dir,
                        task_id=task_id,
                        rel_source_path=rel_path,
                        manifest=manifest,
                    )
            except Exception as e:
                logger.error(f"Error processing {rel_path} for figures: {e}")

        manifest_path = task_figures_dir / "manifest.json"
        manifest.save(manifest_path)
        logger.info(f"Figure extraction complete: {len(manifest.figures)} unique figure assets registered")
        return manifest

    def _cluster_vector_paths(
        self,
        path_objs: List[Any],
        page_w: float,
        page_h: float,
        cluster_dist: float = 18.0,
    ) -> List[Tuple[float, float, float, float]]:
        """Cluster vector paths into discrete figure regions using bounding union."""
        valid_boxes: List[Tuple[float, float, float, float]] = []
        for p in path_objs:
            try:
                left, bottom, right, top = p.get_bounds()
                w = right - left
                h = top - bottom
                # Filter full-page borders / backgrounds
                if w >= page_w * 0.95 and h >= page_h * 0.95:
                    continue
                # Filter full width horizontal separator lines (thin rule)
                if w >= page_w * 0.7 and h < 4:
                    continue
                # Keep significant graphic elements
                if w >= 2 and h >= 2:
                    valid_boxes.append((left, bottom, right, top))
            except Exception:
                pass

        if not valid_boxes:
            return []

        # Merge boxes that are within cluster_dist of each other
        clusters: List[List[float]] = [] # [left, bottom, right, top]
        for box in valid_boxes:
            b_left, b_bottom, b_right, b_top = box
            merged = False
            for c in clusters:
                # Check proximity
                c_left, c_bottom, c_right, c_top = c
                if not (
                    b_right < c_left - cluster_dist
                    or b_left > c_right + cluster_dist
                    or b_top < c_bottom - cluster_dist
                    or b_bottom > c_top + cluster_dist
                ):
                    c[0] = min(c_left, b_left)
                    c[1] = min(c_bottom, b_bottom)
                    c[2] = max(c_right, b_right)
                    c[3] = max(c_top, b_top)
                    merged = True
                    break
            if not merged:
                clusters.append([b_left, b_bottom, b_right, b_top])

        # Second merge pass to coalesce chained clusters
        changed = True
        while changed:
            changed = False
            new_clusters = []
            while clusters:
                c = clusters.pop(0)
                absorbed = False
                for nc in new_clusters:
                    if not (
                        c[2] < nc[0] - cluster_dist
                        or c[0] > nc[2] + cluster_dist
                        or c[3] < nc[1] - cluster_dist
                        or c[1] > nc[3] + cluster_dist
                    ):
                        nc[0] = min(nc[0], c[0])
                        nc[1] = min(nc[1], c[1])
                        nc[2] = max(nc[2], c[2])
                        nc[3] = max(nc[3], c[3])
                        absorbed = True
                        changed = True
                        break
                if not absorbed:
                    new_clusters.append(c)
            clusters = new_clusters

        # Add small padding (e.g. 6 pt) and return tuple
        pad = 6.0
        results = []
        for c in clusters:
            left = max(0.0, c[0] - pad)
            bottom = max(0.0, c[1] - pad)
            right = min(page_w, c[2] + pad)
            top = min(page_h, c[3] + pad)
            results.append((left, bottom, right, top))
        return results

    def _boxes_intersect(
        self,
        b1: Tuple[float, float, float, float],
        b2: Tuple[float, float, float, float],
        threshold: float = 0.5,
    ) -> bool:
        l1, b1_b, r1, t1 = b1
        l2, b2_b, r2, t2 = b2
        ix0 = max(l1, l2)
        iy0 = max(b1_b, b2_b)
        ix1 = min(r1, r2)
        iy1 = min(t1, t2)
        if ix1 <= ix0 or iy1 <= iy0:
            return False
        inter_area = (ix1 - ix0) * (iy1 - iy0)
        area1 = (r1 - l1) * (t1 - b1_b)
        area2 = (r2 - l2) * (t2 - b2_b)
        min_area = min(area1, area2)
        return (inter_area / min_area) >= threshold if min_area > 0 else False

    def _infer_caption(self, text: str, page_num: int) -> str:
        """Infer a caption from neighboring text (e.g. '图 1-2 电路图', '例 3.1')."""
        if not text:
            return f"第 {page_num} 页题图"
        # Look for explicit figure/example caption pattern
        match = re.search(r"(?:图|Figure|Fig\.)\s*([0-9]+[.-][0-9]+|[0-9]+)\s*([^\n\r]{0,30})", text, re.IGNORECASE)
        if match:
            num = match.group(1)
            desc = match.group(2).strip()
            return f"图 {num} {desc}".strip()
        match_ex = re.search(r"(?:例|例题|Example)\s*([0-9]+[.-][0-9]+|[0-9]+)", text, re.IGNORECASE)
        if match_ex:
            return f"例题 {match_ex.group(1)} 配图"
        return f"第 {page_num} 页题图"

    def detect_vector_figure_regions(self, pdf_path: Path, page_index: int = 0) -> List[BoundingBox]:
        """Detect discrete clusters of vector graphics/drawings on a PDF page."""
        doc = pdfium.PdfDocument(pdf_path)
        if page_index < 0 or page_index >= len(doc):
            return []
        page = doc.get_page(page_index)
        page_w, page_h = page.get_size()
        path_objs = [obj for obj in page.get_objects() if obj.type == pdfium.raw.FPDF_PAGEOBJ_PATH]
        if not path_objs:
            return []
        clusters = self._cluster_vector_paths(path_objs, page_w, page_h)
        boxes = []
        for c_left, c_bottom, c_right, c_top in clusters:
            if (c_right - c_left) >= 40 and (c_top - c_bottom) >= 30:
                boxes.append(
                    BoundingBox(
                        x0=max(0.0, min(1.0, c_left / page_w)),
                        y0=max(0.0, min(1.0, (page_h - c_top) / page_h)),
                        x1=max(0.0, min(1.0, c_right / page_w)),
                        y1=max(0.0, min(1.0, (page_h - c_bottom) / page_h)),
                        coord_system=CoordinateSystem.NORMALIZED,
                    )
                )
        return boxes

