"""Comprehensive automated tests for CourseLM Original Source Figures System.

Tests cover:
1. Data models and Manifest management (FigureModels, BoundingBox, Deduplication).
2. Raster image extraction and SHA-256 deduplication.
3. Office documents (DOCX, PPTX) media extraction.
4. PDF embedded figure extraction, vector region clustering, and 300 DPI high-res cropping via pypdfium2.
5. Smart matching (lexical, page, example, figure number) and missing figure scanner.
6. Markdown -> LaTeX -> XeLaTeX compilation pipeline with real graphics.
7. Handout builder figure synchronization and missing asset verification.
8. Studio Workspace figure operations and strict path confinement security.
9. Studio HTTP REST API endpoints for figures.
"""
from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen

import pytest
from PIL import Image, ImageDraw

from modules.figure_extractor import FigureExtractor
from modules.figure_matcher import (
    FigureMatcher,
    audit_handout_figures,
    detect_missing_figures,
)
from modules.figure_models import (
    BoundingBox,
    CoordinateSystem,
    ExtractionType,
    FigureAsset,
    FigureManifest,
    ReviewStatus,
    SourceOccurrence,
)
from modules.handout_builder import _sync_handout_figures, build_handout
from modules.latex_renderer import render_markdown_to_latex, write_latex_document
from studio.workspace import Workspace, confined


# =========================================================================
# 1. Models & Manifest Tests
# =========================================================================

def test_bounding_box_operations():
    """Verify BoundingBox calculations, normalization, and clamping."""
    bbox = BoundingBox(x0=50.0, y0=100.0, x1=250.0, y1=400.0, coord_system=CoordinateSystem.POINTS)
    assert bbox.width == 200.0
    assert bbox.height == 300.0
    assert bbox.area == 60000.0
    assert not bbox.is_empty()

    # Normalization with respect to page dimensions 500x1000
    norm = bbox.normalize(500.0, 1000.0)
    assert norm.x0 == pytest.approx(0.1, abs=0.001)
    assert norm.y0 == pytest.approx(0.1, abs=0.001)
    assert norm.x1 == pytest.approx(0.5, abs=0.001)
    assert norm.y1 == pytest.approx(0.4, abs=0.001)

    # Clamping
    overflow = BoundingBox(x0=-20.0, y0=5.0, x1=600.0, y1=1200.0, coord_system=CoordinateSystem.POINTS)
    clamped = overflow.clamp(500.0, 1000.0)
    assert clamped.x0 == 0.0
    assert clamped.y0 == 5.0
    assert clamped.x1 == 500.0
    assert clamped.y1 == 1000.0


def test_figure_manifest_persistence_and_deduplication(tmp_path):
    """Verify FigureManifest persistence and occurrence-level deduplication."""
    manifest_path = tmp_path / "figures" / "manifest.json"
    manifest = FigureManifest.load(manifest_path, task_id="job_test", course="test_course")
    assert len(manifest.figures) == 0

    asset1 = FigureAsset(
        asset_id="fig_001",
        task_id="job_test",
        file_path="figures/fig_001.png",
        sha256="hash12345678",
        width=800,
        height=600,
        dpi=300,
        caption="基尔霍夫定律图示",
        status=ReviewStatus.PENDING,
        extraction_type=ExtractionType.EMBEDDED,
        occurrences=[
            SourceOccurrence(
                source_file_rel="docs/circuit.pdf",
                source_file_sha256="sourcehash1",
                page_index=1,
            )
        ],
    )
    manifest.add_figure(asset1)
    manifest.save()

    # Reload from disk
    loaded = FigureManifest.load(manifest_path, task_id="job_test", course_id="test_course")
    assert len(loaded.figures) == 1
    assert loaded.figures["fig_001"].caption == "基尔霍夫定律图示"

    # Add duplicate SHA-256 with second occurrence
    asset1_dup = FigureAsset(
        asset_id="fig_dup",
        task_id="job_test",
        file_path="figures/fig_dup.png",
        sha256="hash12345678",
        width=800,
        height=600,
        dpi=300,
        extraction_type=ExtractionType.EMBEDDED,
        occurrences=[
            SourceOccurrence(
                source_file_rel="docs/lecture2.pdf",
                source_file_sha256="sourcehash2",
                page_index=5,
            )
        ],
    )
    existing = loaded.add_figure(asset1_dup)
    assert existing.asset_id == "fig_001"
    assert len(existing.occurrences) == 2
    assert existing.occurrences[1].source_file == "docs/lecture2.pdf"


# =========================================================================
# 2. Standalone Raster Image Extraction Tests
# =========================================================================

def test_extract_standalone_image(tmp_path):
    """Verify scanning and extracting standalone PNG/JPG images."""
    src_dir = tmp_path / "source_docs"
    src_dir.mkdir()
    out_dir = tmp_path / "task_output"
    out_dir.mkdir()

    # Create test image with visible graphics (non-blank)
    img = Image.new("RGB", (300, 200), color="white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([30, 30, 270, 170], outline="black", width=3)
    draw.line([0, 0, 300, 200], fill="blue", width=4)
    img_path = src_dir / "sample_circuit.png"
    img.save(img_path, format="PNG")

    extractor = FigureExtractor(target_dpi=300)
    manifest = FigureManifest.load(out_dir / "figures" / "manifest.json", "task_01", "course_01")
    manifest = extractor.scan_and_extract_directory(src_dir, out_dir, "task_01", manifest)

    assert len(manifest.figures) == 1
    asset = list(manifest.figures.values())[0]
    assert asset.width == 300
    assert asset.height == 200
    assert asset.extraction_type == ExtractionType.STANDALONE
    assert (out_dir / asset.file_path).is_file()


# =========================================================================
# 3. Office Documents (DOCX / PPTX) Extraction Tests
# =========================================================================

def test_extract_docx_and_pptx(tmp_path):
    """Verify embedded-media extraction independently of optional Office renderers."""
    docx_ref = Path("_references/docling/tests/data/docx/sources/docx_grouped_images.docx")
    # The reference architecture slide contains vector groups, not embedded media.
    # Rendering those is a separate capability; don't make this test depend on COM.
    pptx_ref = tmp_path / "synthetic_test.pptx"

    if not docx_ref.is_file():
        import docx
        doc = docx.Document()
        img_p1 = tmp_path / "test1.png"
        im1 = Image.new("RGB", (100, 100), color="white")
        ImageDraw.Draw(im1).line([0, 0, 100, 100], fill="red", width=3)
        im1.save(img_p1)
        doc.add_picture(str(img_p1))
        docx_ref = tmp_path / "synthetic_test.docx"
        doc.save(str(docx_ref))

    if not pptx_ref.is_file():
        import pptx
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        img_p2 = tmp_path / "test2.png"
        im2 = Image.new("RGB", (100, 100), color="white")
        ImageDraw.Draw(im2).rectangle([10, 10, 90, 90], outline="blue", width=3)
        im2.save(img_p2)
        slide.shapes.add_picture(str(img_p2), 0, 0)
        pptx_ref = tmp_path / "synthetic_test.pptx"
        prs.save(str(pptx_ref))

    out_dir = tmp_path / "office_output"
    out_dir.mkdir()
    extractor = FigureExtractor(target_dpi=300)
    manifest = FigureManifest.load(out_dir / "figures" / "manifest.json", "task_office", course="test")

    # DOCX test
    docx_assets = extractor.extract_from_docx(docx_ref, out_dir / "figures", "task_office", "sample.docx", manifest)
    assert len(docx_assets) >= 1
    for a in docx_assets:
        assert (out_dir / a.file_path).is_file()
        assert a.extraction_type == ExtractionType.EMBEDDED

    # PPTX test
    pptx_assets = extractor.extract_from_pptx(pptx_ref, out_dir / "figures", "task_office", "sample.pptx", manifest)
    assert len(pptx_assets) >= 1
    for a in pptx_assets:
        assert (out_dir / a.file_path).is_file()


# =========================================================================
# 4. PDF Extraction, Vector Path Clustering & 300 DPI Cropping
# =========================================================================

def test_pdf_embedded_extraction_and_cropping(tmp_path):
    """Verify PDF embedded extraction, region cropping, and vector clustering."""
    pdf_ref = Path("突击课/电气通信/电路【蜂考突击课】.pdf")
    if not pdf_ref.is_file():
        pytest.skip("Test circuit PDF not found")

    out_dir = tmp_path / "pdf_output"
    out_dir.mkdir()
    figures_dir = out_dir / "figures"
    figures_dir.mkdir()
    manifest = FigureManifest.load(figures_dir / "manifest.json", "task_pdf", "circuit")
    extractor = FigureExtractor(target_dpi=300)

    # 1. Embedded extraction (scan first 3 pages)
    embedded = extractor.extract_from_pdf(pdf_ref, figures_dir, "task_pdf", "circuit.pdf", manifest, max_pages=3)
    assert isinstance(embedded, list)
    for a in embedded:
        assert (out_dir / a.file_path).is_file()
        assert a.sha256

    # 2. Vector cluster detection on page 1
    clusters = extractor.detect_vector_figure_regions(pdf_ref, page_index=1)
    assert isinstance(clusters, list)

    # 3. High-resolution cropping (300 DPI) of a specific region
    crop_bbox = BoundingBox(x0=50.0, y0=100.0, x1=400.0, y1=450.0, coord_system=CoordinateSystem.POINTS)
    cropped_asset = extractor.crop_pdf_region(
        pdf_path=pdf_ref,
        page_index=1,
        bbox=crop_bbox,
        task_figures_dir=figures_dir,
        task_id="task_pdf",
        rel_source_path="circuit.pdf",
        manifest=manifest,
        dpi=300,
        caption="测试电路图裁剪",
        status=ReviewStatus.CONFIRMED,
        extraction_type=ExtractionType.CROPPED,
    )
    assert cropped_asset is not None
    assert cropped_asset.status == ReviewStatus.CONFIRMED
    assert cropped_asset.dpi == 300
    assert (out_dir / cropped_asset.file_path).is_file()
    assert cropped_asset.bbox is not None
    assert 0.0 < cropped_asset.bbox.x0 < 1.0


# =========================================================================
# 5. Figure Matcher & Missing Figure Scanner
# =========================================================================

def test_detect_missing_figures():
    """Verify regex scanning of figure trigger phrases."""
    text_with_missing = (
        "## 考点一：基尔霍夫电压定律\n\n"
        "如图 2-1 所示，沿回路顺时针方向列出电压方程。\n\n"
        "由电路图可知，电阻两端的分压比可直接求得。\n\n"
        "我们知道，电感上的电流不能跃变。"
    )
    warnings = detect_missing_figures(text_with_missing, "chapters/01-基尔霍夫定律.md")
    assert len(warnings) == 2
    assert "图" in warnings[0].matched_text and "所示" in warnings[0].matched_text
    assert warnings[0].file_path == "chapters/01-基尔霍夫定律.md"

    # Text that does NOT require figures
    text_plain = "在电路分析中，我们要掌握画等效电路的方法与基本公式。"
    no_warnings = detect_missing_figures(text_plain, "chapters/02-公式.md")
    assert len(no_warnings) == 0


def test_candidate_matching_and_audit():
    """Verify smart matching priority and comprehensive handout audit."""
    manifest = FigureManifest(task_id="task_match", course="circuit")
    asset = FigureAsset(
        asset_id="fig_match_01",
        task_id="task_match",
        file_path="figures/fig_match_01.png",
        sha256="sha999",
        width=600,
        height=400,
        dpi=300,
        caption="图 2-1 典型单回路电路",
        extraction_type=ExtractionType.EMBEDDED,
        occurrences=[
            SourceOccurrence(
                source_file_rel="circuit.pdf",
                source_file_sha256="pdfsha",
                page_index=3,
                surrounding_text="例题 2.1：如图 2-1 所示，求各支路电流",
            )
        ],
    )
    manifest.add_figure(asset)

    matcher = FigureMatcher(manifest)
    candidates = matcher.find_candidate_figures("如图 2-1 所示，电路中的受控源参数为", limit=3)
    assert len(candidates) >= 1
    best = candidates[0]
    assert best.figure_asset_id == "fig_match_01"
    assert any("图号匹配" in r for r in best.match_reasons)

    # Handout audit test with distant missing figure paragraph
    manuscript = (
        "# 第一章 电路分析\n\n"
        "如图 2-1 所示，该支路电流满足 KCL。\n\n"
        "![图 2-1 典型单回路电路](figures/fig_match_01.png)\n\n"
        "接下来讨论第二小节。\n\n"
        "这一节我们学习叠加定理。\n\n"
        "叠加定理适用于线性含源电路。\n\n"
        "受控源的控制量和受控量必须在同一电路中。\n\n"
        "由电路图可知，另一支路电流为 2A。\n"
    )
    report = audit_handout_figures(manuscript, manifest)
    assert report.total_figures_in_handout == 1
    assert len(report.missing_figure_warnings) >= 1  # "由电路图可知" is missing an image



# =========================================================================
# 6. Markdown -> LaTeX -> XeLaTeX Compilation Pipeline
# =========================================================================

def test_markdown_to_latex_figure_rendering():
    """Verify Markdown image syntax conversion to LaTeX figure environments."""
    md_content = "![电路原理图|width=0.7\\linewidth,align=center](figures/circuit.png)"
    latex = render_markdown_to_latex(md_content)
    assert r"\begin{figure}[H]" in latex
    assert r"\centering" in latex
    assert "width=0.7\\linewidth" in latex
    assert "keepaspectratio" in latex
    assert "figures/circuit.png" in latex
    assert r"\caption{电路原理图}" in latex
    assert r"\end{figure}" in latex


def test_real_xelatex_compilation_with_image(tmp_path):
    """Compile a real XeLaTeX document containing an image to PDF."""
    from main import compile_latex

    # 1. Create a real test PNG
    fig_dir = tmp_path / "figures"
    fig_dir.mkdir()
    img = Image.new("RGB", (400, 300), color=(30, 80, 160))
    d = ImageDraw.Draw(img)
    d.rectangle([50, 50, 350, 250], outline="yellow", width=4)
    img_file = fig_dir / "test_fig.png"
    img.save(img_file)

    # 2. Build LaTeX document
    md_doc = (
        "# 题图测试章节\n\n"
        "下面展示一张经过严格提取的原始电路图：\n\n"
        "![图 1-1 原始电路原理图|width=0.65\\linewidth,align=center](figures/test_fig.png)\n\n"
        "由图可知，该系统稳定运行。"
    )
    tex_path = tmp_path / "main.tex"
    write_latex_document(
        md_doc,
        tex_path,
        "CourseLM 题图系统测试",
        subtitle="原始图像提取与排版验证",
    )

    assert tex_path.is_file()
    tex_src = tex_path.read_text(encoding="utf-8")
    assert "graphicx" in tex_src
    assert "float" in tex_src
    assert "figures/test_fig.png" in tex_src

    # 3. Real XeLaTeX compilation
    ok, log = compile_latex(tex_path)
    assert ok, f"XeLaTeX compilation failed with log:\n{log}"
    pdf_path = tex_path.with_suffix(".pdf")
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 1000  # Non-trivial PDF generated


# =========================================================================
# 7. Handout Builder Figure Synchronization
# =========================================================================

def test_handout_builder_sync_and_missing_check(tmp_path):
    """Unregistered files, including existing ones, must be omitted during build."""
    course_dir = tmp_path / "courses" / "sample_course"
    course_dir.mkdir(parents=True)
    out_dir = tmp_path / "output" / "sample_course"
    md_dir = out_dir / "md"
    ch_dir = md_dir / "chapters"
    figures_dir = out_dir / "figures"
    ch_dir.mkdir(parents=True)
    figures_dir.mkdir(parents=True)

    # Create dummy style config
    (course_dir / "handout.yaml").write_text("mode: sprint\nstyle: academic\ntitle: 测试\n", encoding="utf-8")

    # Create an image in figures_dir
    test_img = figures_dir / "fig_real.png"
    Image.new("RGB", (100, 100), color="blue").save(test_img)

    # Create chapter referencing both existing and missing images
    (md_dir / "00-绪论.md").write_text("# 绪论\n\n课程导读与方法说明", encoding="utf-8")
    ch1 = (
        "# 第一章\n\n"
        "![存在的图片](figures/fig_real.png)\n\n"
        "![不存在的图片](figures/fig_missing.png)\n"
    )
    (ch_dir / "01-test.md").write_text(ch1, encoding="utf-8")

    res = build_handout("sample_course", tmp_path / "output", courses_dir=tmp_path / "courses")
    report = json.loads((out_dir / "handout" / "figure-report.json").read_text(encoding="utf-8"))
    assert report["omitted"] == 2
    assert {item['path'] for item in report['omissions']} == {
        'figures/fig_real.png', 'figures/fig_missing.png'}
    assert '![' not in res.manuscript_path.read_text(encoding='utf-8')
    # File existence alone no longer authorizes copying into the handout.
    handout_fig = out_dir / "handout" / "figures" / "fig_real.png"
    assert not handout_fig.is_file()


# =========================================================================
# 8. Studio Workspace & Path Confinement Security
# =========================================================================

@pytest.fixture
def test_workspace(tmp_path):
    """Create an isolated test workspace with job and course structure."""
    root = tmp_path / "app_root"
    jobid = "a1b2c3d4e5f6"
    course_id = f"physics-{jobid}"
    job_dir = root / "studio_data" / "jobs" / jobid
    course_dir = root / "courses" / course_id
    output_dir = root / "output" / course_id
    src_dir = root / "raw_materials"

    job_dir.mkdir(parents=True)
    course_dir.mkdir(parents=True)
    output_dir.mkdir(parents=True)
    src_dir.mkdir(parents=True)

    # Create source sample image with visible diagram elements
    sample_img = src_dir / "mechanics.png"
    im = Image.new("RGB", (200, 200), color="white")
    d = ImageDraw.Draw(im)
    d.rectangle([20, 20, 180, 180], outline="black", width=3)
    d.line([0, 0, 200, 200], fill="red", width=2)
    im.save(sample_img)

    status = {
        "id": jobid,
        "course": "大学物理",
        "course_id": course_id,
        "status": "draft_ready",
        "stage": "等待审核",
        "percent": 100,
        "artifacts": {},
    }
    request = {
        "course": "大学物理",
        "course_id": course_id,
        "source_path": str(src_dir),
        "mode": "sprint",
        "style": "academic",
    }
    (job_dir / "status.json").write_text(json.dumps(status), encoding="utf-8")
    (job_dir / "request.json").write_text(json.dumps(request), encoding="utf-8")
    (job_dir / "editor.json").write_text("{}", encoding="utf-8")

    ws = Workspace(root, jobid)
    return ws, src_dir, output_dir


def test_workspace_figure_operations_and_confinement(test_workspace):
    """Verify Workspace figure operations and strict confinement against traversal."""
    ws, src_dir, output_dir = test_workspace

    # 1. Scan figures
    figs = ws.scan_figures()
    assert len(figs) >= 1
    asset_id = figs[0]["asset_id"]

    # 2. Update figure
    updated = ws.update_figure(asset_id, caption="修改后的图注", status="confirmed")
    assert updated["caption"] == "修改后的图注"
    assert updated["status"] == "confirmed"

    # 3. Path confinement tests
    with pytest.raises(ValueError, match="文件路径越界|源文件路径格式无效"):
        ws.source_file_path("../../outside.pdf")

    with pytest.raises(ValueError, match="资产编号无效"):
        ws.figure_asset_path("../evil_id")

    with pytest.raises(FileNotFoundError):
        ws.figure_asset_path("fig_000000000000")


# =========================================================================
# 9. Studio HTTP REST API Endpoints
# =========================================================================

@pytest.fixture
def http_server(test_workspace, monkeypatch):
    """Start localhost HTTP server for testing REST endpoints."""
    from studio import server
    ws, _, _ = test_workspace

    monkeypatch.setattr(server, "ROOT", ws.root)
    monkeypatch.setattr(server, "JOBS", ws.job.parent)
    monkeypatch.setattr(server, "ACTIVE", {})

    instance = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    thread = threading.Thread(target=instance.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{instance.server_port}"

    def client(path: str, data: dict | None = None, method: str = "GET"):
        url = base + path
        headers = {"Content-Type": "application/json"}
        body = json.dumps(data).encode("utf-8") if data is not None else None
        req = Request(url, headers=headers, data=body, method=method)
        try:
            with urlopen(req) as resp:
                raw = resp.read()
                content_type = resp.headers.get("Content-Type", "")
                if "application/json" in content_type:
                    return resp.status, json.loads(raw.decode("utf-8"))
                return resp.status, raw
        except HTTPError as err:
            raw = err.read()
            try:
                return err.code, json.loads(raw.decode("utf-8"))
            except Exception:
                return err.code, raw

    yield client, ws
    instance.shutdown()
    instance.server_close()
    thread.join()


def test_http_figure_endpoints(http_server):
    """Verify HTTP figure endpoints: scan, list, asset stream, update, audit."""
    client, ws = http_server
    job_id = ws.job.name
    prefix = f"/api/jobs/{job_id}"

    # 1. Scan figures via POST
    code, scan_res = client(f"{prefix}/figures/scan", data={}, method="POST")
    assert code == 200
    assert isinstance(scan_res, list)
    assert len(scan_res) >= 1

    # 2. List figures via GET
    code, list_res = client(f"{prefix}/figures")
    assert code == 200
    assert isinstance(list_res, list)
    assert len(list_res) >= 1
    asset_id = list_res[0]["asset_id"]

    # 3. Fetch image asset binary stream
    code, asset_bytes = client(f"{prefix}/figures/asset/{asset_id}")
    assert code == 200
    assert len(asset_bytes) > 50

    # 4. Update figure caption and review status
    update_payload = {"asset_id": asset_id, "caption": "HTTP接口更新图注", "status": "confirmed"}
    code, update_res = client(f"{prefix}/figures/update", data=update_payload, method="POST")
    assert code == 200
    assert update_res["caption"] == "HTTP接口更新图注"
    assert update_res["status"] == "confirmed"

    # The picker now uses a versioned question binding, not raw string insertion.
    chapter = ws.output / 'md' / 'chapters' / '01-binding.md'
    chapter.parent.mkdir(parents=True, exist_ok=True)
    chapter.write_text('### 例题 1\n\n已知图示力学系统，求物体所受合力。\n\n**解析**\n\n分析受力。', encoding='utf-8')
    doc = ws.read('chapters/01-binding.md')
    code, bound = client(f'{prefix}/figures/bind', data={
        'path': doc['path'], 'revision': doc['revision'], 'asset_id': asset_id,
        'cursor_line': 3, 'caption': '物体受力示意图'}, method='POST')
    assert code == 200
    assert bound['binding']['task_id'] == job_id
    assert '![物体受力示意图]' in bound['document']['content']

    # 5. Audit figures
    code, audit_res = client(f"{prefix}/figures/audit")
    assert code == 200
    assert "total_figures_in_handout" in audit_res
    assert "missing_figure_warnings" in audit_res

    # 6. Security tests
    # Invalid asset ID
    code, err_res = client(f"{prefix}/figures/asset/../../evil")
    assert code in (400, 404)

    # Invalid job ID
    code, _ = client("/api/jobs/nonexistent123/figures")
    assert code == 404
